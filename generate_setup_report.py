"""Build the Excel report from the CSVs pattern_recognition.py wrote.

    python generate_setup_report.py [--tf M5] [--period 3M] [--pairs EURUSD GBPUSD]

Writes output/reports/<label>.xlsx - Valid Summary, Valid Trades, Invalid Summary,
Invalid Setups, Settings.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import settings as cfg

FILL = PatternFill("solid", fgColor="9999FF")
SIDE = Side(style="thin", color="000000")
BORDER = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)
CENTER = Alignment(horizontal="center", vertical="center")
BOLD = Font(bold=True)
DAILY_LIMIT_DAYS = 45       # more data than this is grouped by month instead of day
TIME_COLUMNS = ["p1_time", "m_time", "p2_time", "entry_time", "exit_time", "be_time"]


def rates(df):
    total = len(df)
    count = lambda name: int((df["result"] == name).sum())
    pct = lambda n: round(100.0 * n / total, 1) if total else 0.0
    return {"Total Trades": total, "Wins": count("Win"), "Losses": count("Loss"),
            "Breakeven": count("Breakeven"),
            "Win Rate %": pct(count("Win")), "Loss Rate %": pct(count("Loss")),
            "Breakeven Rate %": pct(count("Breakeven")),
            "Net Pips": round(df["pnl_pips"].sum(), 1), "Net $": round(df["pnl_cash"].sum(), 2)}


def grouped(df, by, label):
    rows = [{label: str(key), **rates(g)} for key, g in df.groupby(by)]
    return pd.DataFrame(rows, columns=[label, *rates(df)])


def counts(df, column, label):
    out = df.groupby(column).size().sort_values(ascending=False).reset_index()
    out.columns = [label, "Count"]
    return out


def period_table(df):
    times = df["entry_time"].dropna()
    if times.empty:
        return "DAILY SUMMARY", grouped(df, df["entry_time"], "Period")
    if (times.max() - times.min()).days + 1 <= DAILY_LIMIT_DAYS:
        return "DAILY SUMMARY", grouped(df, times.dt.date, "Period")
    return "MONTHLY SUMMARY", grouped(df, times.dt.to_period("M"), "Period")


def as_text(df):
    """Timestamps as plain 2026-09-28 14:20:00 text, blank where missing."""
    df = df.copy()
    for c in df.columns.intersection(TIME_COLUMNS):
        df[c] = df[c].dt.strftime("%Y-%m-%d %H:%M:%S")
    return df


def put(ws, row, col, value, header=False):
    cell = ws.cell(row=row, column=col, value=None if pd.isna(value) else value)
    cell.border, cell.alignment = BORDER, CENTER
    if header:
        cell.font, cell.fill = BOLD, FILL


def table(ws, row, title, df):
    """Merged title bar, header row, body. Returns the next free row."""
    width = len(df.columns)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    for c in range(1, width + 1):
        put(ws, row, c, title if c == 1 else None, header=True)
    for c, name in enumerate(df.columns, 1):
        put(ws, row + 1, c, name, header=True)
    for r, values in enumerate(df.itertuples(index=False), row + 2):
        for c, v in enumerate(values, 1):
            put(ws, r, c, v)
    return row + len(df) + 4


def data_sheet(ws, df):
    for c, name in enumerate(df.columns, 1):
        put(ws, 1, c, name, header=True)
    for r, values in enumerate(df.itertuples(index=False), 2):
        for c, v in enumerate(values, 1):
            put(ws, r, c, v)
    ws.freeze_panes = "A2"


def fit_columns(ws):
    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 3, 11), 70)


def info_table(run):
    be = "the 0 point" if cfg.ZERO_SWING else f"{cfg.BE_TRIGGER_R}R"
    return pd.DataFrame([
        ("Pairs", ", ".join(run.pairs)),
        ("Timeframe / period", f"{run.timeframe} / {run.period}"),
        ("R:R", f"1:{cfg.RR:g}"),
        ("Breakeven", be),
        ("0 point", f"{cfg.ZERO_LOOKBACK_BARS} bars back" if cfg.ZERO_SWING else "off"),
        ("Blackout", f"{cfg.BLACKOUT_START}-{cfg.BLACKOUT_END} broker time"),
        ("Min bars P1 to M", cfg.MIN_BARS_P1_TO_M),
        ("RSI lane", "P1 RSI is a wall" if cfg.P1_RSI_LANE else "70/30 walls only"),
        ("P1 RSI", f"extreme of +/-{cfg.P1_RSI_REACH} candles" if cfg.P1_RSI_PEAK
         else "P1 candle's own RSI"),
    ], columns=["Setting", "Value"])


def main():
    run = cfg.get_run()
    valid = pd.read_csv(run.setups_dir / "valid.csv", parse_dates=TIME_COLUMNS)
    invalid = pd.read_csv(run.setups_dir / "invalid.csv", parse_dates=["p1_time", "m_time", "p2_time"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Valid Summary"
    period_title, period = period_table(valid)
    row = table(ws, 1, "SUMMARY", pd.DataFrame(list(rates(valid).items()), columns=["Metric", "Value"]))
    row = table(ws, row, "PAIR-WISE SUMMARY", grouped(valid, "symbol", "Pair"))
    row = table(ws, row, "PATTERN-WISE SUMMARY", grouped(valid, "pattern", "Pattern"))
    table(ws, row, period_title, period)

    data_sheet(wb.create_sheet("Valid Trades"), as_text(valid))

    ws = wb.create_sheet("Invalid Summary")
    row = table(ws, 1, "SUMMARY", pd.DataFrame([("Total Invalid Setups", len(invalid))], columns=["Metric", "Value"]))
    row = table(ws, row, "PAIR-WISE REJECTIONS", counts(invalid, "symbol", "Pair"))
    row = table(ws, row, "PATTERN-WISE REJECTIONS", counts(invalid, "pattern", "Pattern"))
    table(ws, row, "STAGE-WISE REJECTIONS", counts(invalid, "refused_at", "Stage"))

    data_sheet(wb.create_sheet("Invalid Setups"), as_text(invalid))
    table(wb.create_sheet("Settings"), 1, "SETTINGS", info_table(run))

    for ws in wb:
        fit_columns(ws)
    run.report_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(run.report_path)
    print(f"{len(valid)} valid, {len(invalid)} invalid -> {run.report_path}")


if __name__ == "__main__":
    main()
